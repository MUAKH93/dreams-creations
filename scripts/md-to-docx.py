"""Convert project status markdown to Word (.docx)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "docs" / "DREAMS-CREATIONS-Complete-Status-Deployment-and-MultiTenant.md"
OUT = ROOT / "docs" / "DREAMS-CREATIONS-Complete-Status-Deployment-and-MultiTenant.docx"


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = MD.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []

    def flush_code() -> None:
        nonlocal code_lines
        if code_lines:
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            code_lines = []

    def flush_table() -> None:
        nonlocal table_rows
        if not table_rows:
            return
        cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=cols)
        table.style = "Table Grid"
        for i, row in enumerate(table_rows):
            for j in range(cols):
                text = row[j].strip() if j < len(row) else ""
                table.rows[i].cells[j].text = text
        table_rows = []

    for raw in lines:
        if raw.strip().startswith("```"):
            if in_code:
                in_code = False
                flush_code()
            else:
                flush_table()
                in_code = True
            continue
        if in_code:
            code_lines.append(raw)
            continue
        if raw.startswith("|") and "|" in raw[1:]:
            if re.match(r"^\|[\s\-:|]+\|$", raw.replace(" ", "")):
                continue
            cells = [c.strip() for c in raw.strip("|").split("|")]
            table_rows.append(cells)
            continue
        flush_table()
        if not raw.strip():
            continue
        if raw.startswith("# "):
            doc.add_heading(raw[2:].strip(), level=0)
        elif raw.startswith("## "):
            doc.add_heading(raw[3:].strip(), level=1)
        elif raw.startswith("### "):
            doc.add_heading(raw[4:].strip(), level=2)
        elif raw.startswith("#### "):
            doc.add_heading(raw[5:].strip(), level=3)
        elif raw.startswith("- "):
            doc.add_paragraph(raw[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\.\s", raw):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", raw).strip(), style="List Number")
        else:
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", raw)
            text = re.sub(r"`([^`]+)`", r"\1", text)
            doc.add_paragraph(text)

    flush_code()
    flush_table()
    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
